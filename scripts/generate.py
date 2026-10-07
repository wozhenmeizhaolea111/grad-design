#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""generate.py — 毕业设计文档生成器。

子命令:
    task     从 JSON 输入填充“毕业设计任务书”（templates/renwushu.docx）
    outline  按专业大类+选题类别生成正文骨架 .docx（含封面、章节标题、撰写提示、
             参考文献占位，全部按 references/format.md 规范排版）

用法:
    python generate.py task --input 任务书.json --output 任务书-已填.docx
    python generate.py outline --major electronic --category product \
        --title "基于单片机的XX系统设计与制作" --output 正文骨架.docx
"""
import argparse
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
TASK_TEMPLATE = os.path.join(HERE, "..", "templates", "renwushu.docx")

FORM_OPTIONS = ["方案", "物化产品", "文化作品", "软件", "其他"]

# (major, category) -> [(标题, 撰写提示), ...]，参考文献固定追加
SECTION_PLAN = {
    ("electronic", "product"): [
        ("一、引言", "课题背景与来源、国内外现状、研究目的与意义；与任务书“选题目的意义”呼应。"),
        ("二、需求（功能）分析", "产品功能需求、性能技术指标、使用场景与约束（成本、环保、安全）；逐条可度量。"),
        ("三、设计方案分析和拟定", "至少 2 个候选方案对比（原理、器件/技术选型、优缺点），论证并确定最终方案。"),
        ("四、技术参数确定", "关键参数计算、元器件选型依据、程序流程/算法设计；公式与数据准确、来源可查。"),
        ("五、设计与实现过程", "原理图、PCB 图、装配图、程序流程图、程序清单、元器件清单；图文并茂记录实现与调试过程。"),
        ("六、功能与效果分析", "测试方法与数据、功能/技术指标达成情况、与需求对照；成本、环保、安全合规说明。"),
        ("七、结论", "总结成果、指出不足与改进方向。"),
    ],
    ("electronic", "scheme"): [
        ("一、引言", "项目背景、来源（企业真实项目优先）、要解决的实际问题。"),
        ("二、需求分析", "业务/技术需求、现状问题、约束条件与设计目标。"),
        ("三、设计方案分析和拟定", "多方案对比论证（技术路线、标准运用、资源与成本），确定推荐方案。"),
        ("四、技术路线与参数确定", "技术路线/拓扑结构、关键参数与计算、标准规范引用；数据详实充分。"),
        ("五、方案设计过程", "系统结构图、网络拓扑图、工程图纸、配置/工程表单；过程性结论逐节给出。"),
        ("六、预期效果分析", "方案实施后的预期效果、可操作性、效益与风险分析。"),
        ("七、结论", "总结方案价值、局限与后续建议。"),
    ],
    ("electronic", "process"): [
        ("一、引言", "产品对象、工艺现状问题、工艺设计与实施的目标。"),
        ("二、设计需求分析", "技术要求、产量/质量/成本约束、安全环保要求。"),
        ("三、设计方案分析和拟定", "工艺方案（路线/工具/设备）比较论证，确定工艺方案。"),
        ("四、工艺路线与工序设计", "工艺流程图、工艺过程卡、工序卡；工序参数、工具与检验要求填写完整。"),
        ("五、技术参数确定", "工艺参数选取与计算依据、标准规范引用；数据合理准确。"),
        ("六、工艺设计效果分析", "工艺验证/试产结果、质量与效率对比、照片或视频佐证说明。"),
        ("七、结论", "总结工艺方案的可行性与改进点。"),
    ],
    ("arts", "work"): [
        ("一、引言", "课题背景与来源、设计定位要解决的问题、目标用户/场景。"),
        ("二、设计定位与需求分析", "需求调研与信息检索过程、设计定位（原创性、实用性、审美性目标）。"),
        ("三、构思与创意设计", "创意构思、方案比选、元素/技法/材料/工艺选择依据；体现艺术表达。"),
        ("四、设计与制作过程", "设计稿迭代、制作工艺流程、关键节点与调整记录；过程完整可追溯。"),
        ("五、成果特点与效果分析", "作品照片/视频呈现，艺术性与功能性分析，与设计定位对照；应用价值。"),
        ("六、结论", "总结设计思路与成果特点、不足与展望。"),
    ],
    ("arts", "scheme"): [
        ("一、引言", "项目背景、真实项目来源、要解决的实际问题。"),
        ("二、项目背景与问题分析", "现状调研、问题梳理、需求与约束（成本、环保、创新、安全）。"),
        ("三、工作思路与创意方案依据", "设计依据、创意来源、多方案比较与定位论证。"),
        ("四、方案设计与实施过程", "完整方案内容：图纸（平面图/施工图/效果图/设计说明）、实施步骤与进度。"),
        ("五、预期效果分析", "实施后预期效果、可操作性、行业标准符合性。"),
        ("六、结论", "总结方案价值、局限与建议。"),
    ],
}

# 通用骨架（templates/正文模板.docx；类别章节由撰写者按 deliverables.md 展开）
SECTION_PLAN[("basic", "basic")] = [
    ("一、引言", "课题背景与来源、国内外现状、研究目的与意义；说明要解决的实际问题。"),
    ("二、正文", "按所属选题类别的章节结构展开（见 references/deliverables.md）："
     "需求/功能分析 → 方案分析和拟定 → 技术参数确定 → 设计实现过程 → 效果分析；"
     "配图配表并逐节给出阶段性小结。"),
    ("三、结论", "总结全部工作与成果、指出不足与改进方向。"),
]


def set_run(run, font="宋体", size=12, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    return run


def set_indent(p, chars, size_pt):
    p.paragraph_format.first_line_indent = Pt(size_pt * chars)
    ind = p._p.pPr.find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), str(int(chars * 100)))


def add_para(doc, text, font="宋体", size=12, bold=False, align=None,
             indent_chars=None, line_spacing=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if line_spacing is not None:
        p.paragraph_format.line_spacing = line_spacing
    if indent_chars is not None:
        set_indent(p, indent_chars, size)
    if text:
        set_run(p.add_run(text), font=font, size=size, bold=bold)
    return p


def add_field(p, instr, font="宋体", size=10.5):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None),
                      ("text", "1"), ("end", None)):
        r = p.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {txt} "
            r._r.append(el)
        elif kind == "text":
            r.text = txt
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        set_run(r, font=font, size=size)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    footer = sec.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("第"), size=10.5)
    add_field(p, "PAGE")
    set_run(p.add_run("页，共"), size=10.5)
    add_field(p, "NUMPAGES")
    set_run(p.add_run("页"), size=10.5)
    return doc


# ---------------------------------------------------------------- outline ----

def build_outline(args):
    if args.category == "basic":
        key = ("basic", "basic")
    else:
        key = (args.major, args.category)
    if key not in SECTION_PLAN:
        sys.exit(f"error: --major {args.major} 不支持 --category {args.category}")
    sections = SECTION_PLAN[key]

    doc = setup_page(Document())

    # 封面
    add_para(doc, "", line_spacing=None)
    add_para(doc, "", line_spacing=None)
    add_para(doc, args.title, font="黑体", size=22,
             align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=None)
    add_para(doc, "", line_spacing=None)
    cover = [("姓　　名", args.student_name), ("学　　号", args.student_id),
             ("专业名称", args.program), ("班　　级", args.class_name),
             ("指导教师", args.advisor)]
    for label, val in cover:
        add_para(doc, f"{label}　　{val or ''}", size=16,
                 align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=None)
    # 分页
    br = doc.add_paragraph()
    br.add_run().add_break(WD_BREAK.PAGE)

    # 正文标题：上下各空一行
    add_para(doc, "", line_spacing=None)
    add_para(doc, args.title, font="黑体", size=22,
             align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=None)
    add_para(doc, "", line_spacing=None)

    for heading, prompt in sections:
        add_para(doc, heading, size=14, bold=True, indent_chars=2)
        add_para(doc, f"【撰写提示】{prompt}", indent_chars=2)

    # 参考文献：小四黑体顶格，标题前空两行
    add_para(doc, "", line_spacing=None)
    add_para(doc, "", line_spacing=None)
    add_para(doc, "参考文献", font="黑体", size=12, bold=False,
             line_spacing=1.5)
    for i, ref in enumerate([
        "作者，文题（书名），出版地，出版社，出版年，版次",
        "作者，文献题名[J]，刊名，出版年，卷(期)：起止页码",
        "作者，电子文献题名[EB/OL]，发布日期，获取路径",
    ], 1):
        add_para(doc, f"[{i}] {ref}", font="仿宋", size=10.5, line_spacing=1.5)

    doc.save(args.output)
    print(f"OK 已生成正文骨架: {args.output}")
    print(f"   大类={args.major} 类别={args.category} 章节数={len(sections)}")


# ------------------------------------------------------------------- task ----

def logical_cells(row):
    """按 tc 元素去重，返回行内逻辑单元格列表。"""
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c)
    return out


def set_cell(cell, text):
    """用 text 覆盖单元格内容（保留首个 run 的字体格式）。"""
    first = cell.paragraphs[0]
    # 删除多余段落
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if first.runs:
        first.runs[0].text = text
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.add_run(text)


def fill_task(args):
    if not os.path.exists(args.input):
        sys.exit(f"error: 找不到输入文件 {args.input}")
    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    tpl = args.template or TASK_TEMPLATE
    if not os.path.exists(tpl):
        sys.exit(f"error: 找不到任务书模板 {tpl}")
    doc = Document(tpl)
    t = doc.tables[0]
    rows = t.rows

    def cell_text(row_idx):
        return logical_cells(rows[row_idx])

    # r0: 二级学院 | 值 | 专业 | 值
    c0 = cell_text(0)
    if len(c0) >= 4 and data.get("college"):
        set_cell(c0[1], data["college"])
    if len(c0) >= 4 and data.get("program"):
        set_cell(c0[3], data["program"])
    # r1: 班级 | 值 | 学号 | 值 | 指导教师 | 值
    c1 = cell_text(1)
    if len(c1) >= 6:
        if data.get("class_name"):
            set_cell(c1[1], data["class_name"])
        if data.get("student_id"):
            set_cell(c1[3], data["student_id"])
        if data.get("advisor"):
            set_cell(c1[5], data["advisor"])
    # r2: 题目
    c2 = cell_text(2)
    if len(c2) >= 2 and data.get("title"):
        set_cell(c2[1], data["title"])
    # r3: 时间
    c3 = cell_text(3)
    if len(c3) >= 2 and data.get("start_date"):
        end = data.get("end_date", "")
        set_cell(c3[1], f"{data['start_date']} 至 {end}" if end else data["start_date"])
    # r4: 形式（勾选）
    c4 = cell_text(4)
    if len(c4) >= 2 and data.get("form"):
        form = data["form"]
        if form not in FORM_OPTIONS:
            sys.exit(f"error: form 必须是 {FORM_OPTIONS} 之一，收到: {form}")
        line = "  ".join(("☑" if o == form else "☐") + o for o in FORM_OPTIONS)
        set_cell(c4[1], line)
    # r5: 选题目的意义及主要任务要求
    c5 = cell_text(5)
    if data.get("purpose") and c5:
        cell = c5[0]
        target = None
        seen_label = False
        for p in cell.paragraphs:
            if not seen_label:
                if "选题目的意义" in p.text:
                    seen_label = True
                continue
            if not p.text.strip():
                target = p
                break
        if target is None:
            target = cell.add_paragraph()
        for r in list(target.runs):
            r._element.getparent().remove(r._element)
        target.add_run(data["purpose"])
    # r6-r9: 四周计划（逻辑单元格: 计划安排 | 第N周 | 内容）
    weeks = data.get("weeks") or []
    for i, row_idx in enumerate((6, 7, 8, 9)):
        if i >= len(weeks) or not weeks[i]:
            continue
        cells = cell_text(row_idx)
        if len(cells) >= 3:
            set_cell(cells[2], str(weeks[i]))

    doc.save(args.output)
    print(f"OK 任务书已填充: {args.output}")


def main():
    ap = argparse.ArgumentParser(description="毕业设计文档生成器")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p1 = sub.add_parser("task", help="填充毕业设计任务书")
    p1.add_argument("--input", required=True, help="JSON 输入文件")
    p1.add_argument("--output", required=True, help="输出 .docx 路径")
    p1.add_argument("--template", help="自定义任务书模板路径")

    p2 = sub.add_parser("outline", help="生成正文骨架")
    p2.add_argument("--major", required=True, choices=["electronic", "arts"],
                    help="electronic=电子信息大类 / arts=文化艺术大类")
    p2.add_argument("--category", required=True,
                    help="electronic: product|scheme|process; arts: work|scheme")
    p2.add_argument("--title", required=True, help="毕业设计题目")
    p2.add_argument("--output", required=True, help="输出 .docx 路径")
    p2.add_argument("--student-name", default="")
    p2.add_argument("--student-id", default="")
    p2.add_argument("--program", default="")
    p2.add_argument("--class-name", default="")
    p2.add_argument("--advisor", default="")

    args = ap.parse_args()
    if args.cmd == "task":
        fill_task(args)
    else:
        build_outline(args)


if __name__ == "__main__":
    main()
