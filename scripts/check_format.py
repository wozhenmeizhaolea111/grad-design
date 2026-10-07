#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""check_format.py — 咸阳职业技术学院毕业设计排版规范检查。

依据 references/format.md（《咸阳职业技术学院毕业设计规范格式》）逐项检查
.docx 文件，输出 [OK]/[FAIL]/[WARN] 报告。存在 FAIL 时退出码为 1。

用法:
    python check_format.py 正文.docx [--json]
"""
import argparse
import json
import re
import sys

from docx import Document
from docx.oxml.ns import qn

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# 字号（磅） -> 中文称谓
SIZE_CN = {22: "二号", 16: "三号", 14: "四号", 12: "小四", 10.5: "五号", 9: "小五"}
BODY_FONT_OK = ("宋体", "SimSun")
HEI_FONT_OK = ("黑体", "SimHei")
FANG_FONT_OK = ("仿宋", "FangSong", "FangSong_GB2312")
TABLE_FONT_OK = ("宋体", "SimSun", "黑体", "SimHei")
HALF_PUNCT_AFTER_CJK = re.compile(r"[\u4e00-\u9fff][,.;:?!]")
HALF_PUNCT_BEFORE_CJK = re.compile(r"[,.;:?!][\u4e00-\u9fff]")


def _in(run, names):
    return any((n or "").lower() == x.lower() for n in [run] for x in names if n)


def run_east_asia(run):
    rPr = run._element.rPr
    if rPr is None or rPr.rFonts is None:
        return None
    return rPr.rFonts.get(qn("w:eastAsia"))


def run_size_pt(run):
    if run.font.size is not None:
        return round(run.font.size.pt, 1)
    return None


def para_first_run(p):
    for r in p.runs:
        if r.text.strip():
            return r
    return p.runs[0] if p.runs else None


def _ind_chars(pPr, size_pt):
    """从 w:ind 解析首行缩进（字符数）。"""
    if pPr is None:
        return None
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return None
    flc = ind.get(qn("w:firstLineChars"))
    if flc is not None:
        try:
            return int(flc) / 100.0
        except ValueError:
            pass
    fl = ind.get(qn("w:firstLine"))
    if fl is not None and size_pt:
        try:
            return round(int(fl) / 20.0 / size_pt, 2)
        except ValueError:
            pass
    return 0.0 if (flc is None and fl is not None) or ind.get(qn("w:hanging")) else None


def first_line_chars(p, size_pt):
    """沿 段落 → 样式链 + 文档默认值 解析首行缩进字符数。"""
    v = _ind_chars(p._p.pPr if p._p.pPr is not None else None, size_pt)
    if v is not None:
        return v
    style = p.style
    while style is not None:
        pPr = getattr(style.element, "pPr", None)
        v = _ind_chars(pPr, size_pt)
        if v is not None:
            return v
        style = style.base_style
    docPr = p.part.document.styles.element.find(qn("w:docDefaults"))
    if docPr is not None:
        pPrD = docPr.find(qn("w:pPrDefault"))
        if pPrD is not None:
            return _ind_chars(pPrD.find(qn("w:pPr")), size_pt)
    return None


def _spacing(pPr):
    if pPr is None:
        return None
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        return None
    line, rule = sp.get(qn("w:line")), sp.get(qn("w:lineRule"))
    if line and rule == "auto":
        try:
            return round(int(line) / 240.0, 2)
        except ValueError:
            return None
    return None


def line_spacing(p):
    sp = _spacing(p._p.pPr if p._p.pPr is not None else None)
    if sp is not None:
        return sp
    style = p.style
    while style is not None:
        sp = _spacing(getattr(style.element, "pPr", None))
        if sp is not None:
            return sp
        style = style.base_style
    # docDefaults
    docPr = p.part.document.styles.element.find(qn("w:docDefaults"))
    if docPr is not None:
        pPrD = docPr.find(qn("w:pPrDefault"))
        if pPrD is not None:
            return _spacing(pPrD.find(qn("w:pPr")))
    return None


def para_font(p):
    """段落主体字体（取首个非空 run）。"""
    r = para_first_run(p)
    if r is None:
        return None, None
    return run_east_asia(r), run_size_pt(r)


def para_aligned_center(p):
    return p.alignment is not None and int(p.alignment) == 1


def para_aligned_left(p):
    # None = 继承样式（通常左对齐），视为通过
    return p.alignment is None or int(p.alignment) == 0


def footer_has_page_fields(section):
    xml = section.footer._element.xml if section.footer is not None else ""
    return ("PAGE" in xml) and ("NUMPAGES" in xml)


def footer_centered(section):
    for p in section.footer.paragraphs:
        if p.alignment is not None and int(p.alignment) == 1:
            return True
    return False


def footer_font_ok(section):
    seen = False
    for p in section.footer.paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            seen = True
            ea, sz = run_east_asia(r), run_size_pt(r)
            if sz not in (None, 10.5):
                return False, seen
            if ea is not None and not any(x in ea for x in BODY_FONT_OK):
                return False, seen
    return True, seen


def main():
    ap = argparse.ArgumentParser(description="毕业设计排版规范检查")
    ap.add_argument("file", help="待检查的 .docx")
    ap.add_argument("--json", action="store_true", help="输出 JSON 报告")
    args = ap.parse_args()

    errors, warnings, oks = [], [], []

    def err(code, msg, count=None, examples=None):
        s = f"{code} {msg}"
        if count is not None:
            s += f": {count} 处"
        if examples:
            s += f" （如: {'; '.join(examples[:3])}）"
        errors.append(s)

    def warn(code, msg, count=None, examples=None):
        s = f"{code} {msg}"
        if count is not None:
            s += f": {count} 处"
        if examples:
            s += f" （如: {'; '.join(examples[:3])}）"
        warnings.append(s)

    def ok(msg):
        oks.append(msg)

    doc = Document(args.file)

    # ---- 表单自动识别：任务书/指导记录/评分表只查页面 ----
    # 表格内标记（任务书/指导记录），或“无缩进且较短”的标题段落（文档名）
    def _is_form_title(p):
        if len(p.text.strip()) > 40:
            return False
        if first_line_chars(p, 12) not in (None, 0.0):
            return False
        return any(m in p.text for m in ("任务书", "指导记录", "评分表", "成绩评定表"))

    is_form = any(
        any(m in c.text for m in ("毕业设计题目", "指导主要内容记录"))
        for t in doc.tables for row in t.rows for c in row.cells
    ) or any(_is_form_title(p) for p in doc.paragraphs)

    # ---- 1. 页面规格 A4 ----
    bad_pages = []
    for i, sec in enumerate(doc.sections):
        w, h = round(sec.page_width.mm), round(sec.page_height.mm)
        if abs(w - 210) > 1 or abs(h - 297) > 1:
            bad_pages.append(f"节{i+1}: {w}x{h}mm")
    if bad_pages:
        err("E01", "页面非 A4(210x297mm)", examples=bad_pages)
    else:
        ok("E01 页面 A4")

    if is_form:
        ok("表单文档（任务书/指导记录/评分表）：跳过正文排版检查")
        report = {"file": args.file, "doc_type": "form", "ok": oks,
                  "errors": errors, "warnings": warnings, "passed": not errors}
        if args.json:
            print(json.dumps(report, ensure_ascii=False, indent=2))
        else:
            for s in oks:
                print(f"[OK]   {s}")
            for s in errors:
                print(f"[FAIL] {s}")
            for s in warnings:
                print(f"[WARN] {s}")
            print("-" * 60)
            if errors:
                print(f"结果: 不通过 — {len(errors)} 项错误, {len(warnings)} 项警告")
            else:
                print(f"结果: 通过 — 0 项错误, {len(warnings)} 项警告")
        sys.exit(1 if errors else 0)

    # ---- 2. 页码 第X页，共X页 ----
    foot_bad = []
    for i, sec in enumerate(doc.sections):
        if not footer_has_page_fields(sec):
            foot_bad.append(f"节{i+1}: 缺 PAGE/NUMPAGES 域")
        elif not footer_centered(sec):
            foot_bad.append(f"节{i+1}: 页码未居中")
        else:
            fok, seen = footer_font_ok(sec)
            if seen and not fok:
                foot_bad.append(f"节{i+1}: 页码字体应为五号宋体")
    if foot_bad:
        err("E02", "页码不合规(第X页，共X页/五号宋体/居中)", examples=foot_bad)
    else:
        ok("E02 页码 第X页，共X页（五号宋体居中）")

    # ---- 段落分类 ----
    titles, headings, bodies, refs_body, cover_lines, unknown = [], [], [], [], [], []
    for p in doc.paragraphs:
        text = p.text.strip()
        ea, sz = para_font(p)
        if not text:
            continue
        if sz == 22:
            titles.append(p)
        elif sz == 14:
            headings.append(p)
        elif sz == 12:
            if text == "参考文献":
                refs_body.append(("HEAD", p))
            elif refs_body:
                refs_body.append(("ENTRY", p))
            else:
                bodies.append(p)
        elif sz == 16:
            cover_lines.append(p)
        elif sz == 10.5 and refs_body:
            refs_body.append(("ENTRY", p))
        else:
            unknown.append(p)

    # ---- 3-5. 正文（小四宋体 / 首行缩进2字符 / 1.5倍行距） ----
    bad_font, bad_indent, bad_space, unresolved = [], [], [], []
    for p in bodies:
        ea, sz = para_font(p)
        if ea is None or not any(x in ea for x in BODY_FONT_OK):
            bad_font.append(p.text.strip()[:20])
        v = first_line_chars(p, 12)
        if v is None:
            unresolved.append(p.text.strip()[:20])
        elif abs(v - 2.0) > 0.01:
            bad_indent.append(f"{p.text.strip()[:12]}(缩进{v}字符)")
        sp = line_spacing(p)
        if sp is None:
            unresolved.append(p.text.strip()[:20])
        elif abs(sp - 1.5) > 0.01:
            bad_space.append(f"{p.text.strip()[:12]}({sp}倍)")
    if bodies:
        if bad_font:
            err("E03", "正文应为小四宋体", count=len(bad_font), examples=bad_font)
        else:
            ok(f"E03 正文字体小四宋体（{len(bodies)}段）")
        if bad_indent:
            err("E04", "正文首行应缩进2字符", count=len(bad_indent), examples=bad_indent)
        else:
            ok("E04 正文首行缩进2字符")
        if bad_space:
            err("E05", "正文行距应为1.5倍", count=len(bad_space), examples=bad_space)
        else:
            ok("E05 正文1.5倍行距")
    else:
        warn("W00", "未识别到正文段落（小四段落为空），跳过正文检查")
    if unresolved:
        warn("W15", "部分段落的缩进/行距依赖样式默认值，未能直接解析",
             count=len(unresolved), examples=list(dict.fromkeys(unresolved))[:3])

    # ---- 6. 标题 二号黑体居中 ----
    if not titles:
        err("E06", "缺少主标题（二号黑体居中，上下各空一行）")
    else:
        t = titles[0]
        ea, _ = para_font(t)
        probs = []
        if ea is None or not any(x in ea for x in HEI_FONT_OK):
            probs.append("非黑体")
        if not para_aligned_center(t):
            probs.append("未居中")
        if probs:
            err("E06", "主标题应为二号黑体居中", examples=probs)
        else:
            ok("E06 主标题二号黑体居中")

    # ---- 7. 文中标题（4号宋体加粗，左对齐缩进2字符，末尾无标点） ----
    bad_h = []
    for p in headings:
        text = p.text.strip()
        ea, _ = para_font(p)
        probs = []
        if ea is None or not any(x in ea for x in BODY_FONT_OK):
            probs.append("非宋体")
        if not p.runs or not any(r.font.bold for r in p.runs):
            probs.append("未加粗")
        if not para_aligned_left(p):
            probs.append("非左对齐")
        v = first_line_chars(p, 14)
        if v is None or abs(v - 2.0) > 0.01:
            probs.append("未缩进2字符")
        if text and text[-1] in "。．，、；：？！":
            probs.append("末尾带标点")
        if probs:
            bad_h.append(f"{text[:16]}({'/'.join(probs)})")
    if headings:
        if bad_h:
            err("E07", "文中标题应为4号宋体加粗、左对齐缩进2字符、末尾不加标点",
                count=len(bad_h), examples=bad_h)
        else:
            ok(f"E07 文中标题 4号宋体加粗（{len(headings)}个）")

    # ---- 8-9. 参考文献 ----
    head_entries = [x for x in refs_body if x[0] == "HEAD"]
    entries = [x for x in refs_body if x[0] == "ENTRY"]
    if not head_entries:
        err("E08", "缺少“参考文献”标题")
    else:
        p = head_entries[0][1]
        ea, sz = para_font(p)
        probs = []
        if ea is None or not any(x in ea for x in HEI_FONT_OK):
            probs.append("非黑体")
        if sz != 12:
            probs.append("非小四")
        # 空两行：前两个段落应为空
        paras = doc.paragraphs
        idx = next((i for i, q in enumerate(paras) if q._p is p._p), None)
        blanks = 0
        if idx is not None:
            for q in reversed(paras[:idx]):
                if q.text.strip():
                    break
                blanks += 1
        else:
            blanks = 2
        if blanks < 2:
            probs.append(f"标题前仅空{blanks}行(应空两行)")
        if first_line_chars(p, 12) not in (None, 0.0):
            probs.append("未顶格")
        if probs:
            err("E08", "“参考文献”应为小四黑体、空两行、顶格", examples=probs)
        else:
            ok("E08 参考文献标题 小四黑体顶格（前空两行）")
    if entries:
        bad_ref = []
        for i, (_, p) in enumerate(entries, 1):
            text = p.text.strip()
            ea, sz = para_font(p)
            probs = []
            if not re.match(r"^\[\d+\]", text):
                probs.append("序号非[1]形式")
            if sz != 10.5:
                probs.append("非五号")
            if ea is None or not any(x in ea for x in FANG_FONT_OK):
                probs.append("非仿宋")
            v = first_line_chars(p, 10.5)
            if v not in (None, 0.0):
                probs.append("未顶格")
            if probs:
                bad_ref.append(f"{text[:16]}({'/'.join(probs)})")
        if bad_ref:
            err("E09", "参考文献条目应为[序号]五号仿宋顶格", count=len(bad_ref), examples=bad_ref)
        else:
            ok(f"E09 参考文献条目 [序号]五号仿宋顶格（{len(entries)}条）")
    elif head_entries:
        warn("W09", "“参考文献”标题下未识别到条目")

    # ---- 10. 表格（小五；表号/标题 小五黑体） ----
    if doc.tables:
        bad_cell_font, bad_cell_size = [], []
        for ti, t in enumerate(doc.tables, 1):
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if not p.text.strip():
                            continue
                        ea, sz = para_font(p)
                        if sz != 9:
                            bad_cell_size.append(f"表{ti}:{p.text.strip()[:10]}")
                        elif ea is not None and not any(x in ea for x in TABLE_FONT_OK):
                            bad_cell_font.append(f"表{ti}:{p.text.strip()[:10]}")
        if bad_cell_size:
            err("E10", "表格内文字及数码应为小五号", count=len(bad_cell_size),
                examples=bad_cell_size)
        else:
            ok(f"E10 表内文字小五号（{len(doc.tables)}张表）")
        if bad_cell_font:
            err("E10b", "表格内文字应为小五宋体", count=len(bad_cell_font),
                examples=bad_cell_font)
        # 表号连续性
        cap_nums = []
        for p in doc.paragraphs:
            m = re.match(r"^表\s*(\d+)", p.text.strip())
            if m:
                cap_nums.append(int(m.group(1)))
        if cap_nums and cap_nums != list(range(1, len(cap_nums) + 1)):
            warn("W10", f"表号未按 1..n 连续编号: {cap_nums}")
        elif cap_nums:
            ok("W10 表号连续编号")
    else:
        ok("E10 无表格（跳过表格检查）")

    # ---- 11. 插图 ----
    n_images = len(doc.inline_shapes)
    fig_caps = []
    for p in doc.paragraphs:
        m = re.match(r"^图\s*(\d+)", p.text.strip())
        if m:
            fig_caps.append((int(m.group(1)), p))
    if n_images:
        if not fig_caps:
            err("E11", f"存在 {n_images} 个插图但缺“图1、图2…”图序图名")
        else:
            bad_fig = []
            for n, p in fig_caps:
                ea, sz = para_font(p)
                if sz != 9 or ea is None or not any(x in ea for x in HEI_FONT_OK):
                    bad_fig.append(p.text.strip()[:20])
            if bad_fig:
                err("E11b", "图序及图名应为小5号黑体", count=len(bad_fig), examples=bad_fig)
            else:
                ok(f"E11 图序图名小5黑体（{len(fig_caps)}个）")
            nums = [n for n, _ in fig_caps]
            if nums != list(range(1, len(nums) + 1)):
                warn("W11", f"图序未连续编号: {nums}")
    else:
        ok("E11 无插图（跳过插图检查）")

    # ---- 封面信息行（三号宋体）—— 警告级 ----
    if cover_lines:
        bad_cover = [p.text.strip()[:14] for p in cover_lines
                     if (para_font(p)[0] is None
                         or not any(x in para_font(p)[0] for x in BODY_FONT_OK))]
        if bad_cover:
            warn("W12", "封面信息行应为三号宋体", count=len(bad_cover), examples=bad_cover)
        else:
            ok("W12 封面信息行三号宋体")

    # ---- 中英文标点混用 —— 警告级 ----
    punct = []
    for p in bodies + [x[1] for x in entries]:
        t = p.text
        if HALF_PUNCT_AFTER_CJK.search(t) or HALF_PUNCT_BEFORE_CJK.search(t):
            punct.append(t.strip()[:20])
    if punct:
        warn("W13", "疑似中文语境误用半角标点", count=len(punct), examples=punct)
    else:
        ok("W13 中英文标点使用正常")

    # ---- 未分类段落 —— 提示 ----
    if unknown:
        warn("W14", "存在字号非规范值的段落（自行核对）", count=len(unknown),
             examples=[p.text.strip()[:16] for p in unknown])

    # ---- 报告 ----
    report = {"file": args.file, "ok": oks, "errors": errors, "warnings": warnings,
              "passed": not errors}
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        for s in oks:
            print(f"[OK]   {s}")
        for s in errors:
            print(f"[FAIL] {s}")
        for s in warnings:
            print(f"[WARN] {s}")
        print("-" * 60)
        if errors:
            print(f"结果: 不通过 — {len(errors)} 项错误, {len(warnings)} 项警告")
        else:
            print(f"结果: 通过 — 0 项错误, {len(warnings)} 项警告")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
